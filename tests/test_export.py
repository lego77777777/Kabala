import csv
from datetime import date, datetime

import pytest
from openpyxl import load_workbook

from kabala import i18n
from kabala.autoscan import is_due, runs_at_login, set_run_at_login
from kabala.export import export_csv, export_xlsx, filter_period, period_range
from kabala.models import Receipt


def r(company, day, amount, currency="ILS", path=""):
    return Receipt(id=day + company, message_id=f"<{day}{company}>", company=company, date=day, amount=amount,
                   currency=currency, subject=f"קבלה {company}", sender="x@y.com", confidence=0.95,
                   invoice_number="123", file_path=path)


@pytest.fixture
def receipts(tmp_path):
    pdf = tmp_path / "Receipts" / "Netflix" / "2026" / "2026-10" / "2026-10-14_Netflix_54.90ILS.pdf"
    pdf.parent.mkdir(parents=True)
    pdf.write_bytes(b"%PDF")
    return [
        r("Netflix", "2026-10-14", 54.90, path=str(pdf)),
        r("בזק", "2026-10-03", 129.00),
        r("Amazon", "2026-09-20", 49.20, "USD"),
        r("Netflix", "2026-09-14", 54.90),
        r("חברת החשמל", "2025-12-01", 412.37),
    ]


@pytest.fixture(autouse=True)
def hebrew():
    i18n.set_language("he")
    yield
    i18n.set_language("he")


def test_periods():
    today = date(2026, 10, 20)
    assert period_range("this_month", today) == (date(2026, 10, 1), today)
    assert period_range("last_month", today) == (date(2026, 9, 1), date(2026, 9, 30))
    assert period_range("last_year", today) == (date(2025, 1, 1), date(2025, 12, 31))
    assert period_range("all", today) is None
    assert period_range("last_month", date(2026, 1, 5)) == (date(2025, 12, 1), date(2025, 12, 31))


def test_filter_period(receipts):
    rows = filter_period(receipts, "last_month", date(2026, 10, 20))
    assert [x.company for x in rows] == ["Netflix", "Amazon"]  # sorted by date


def test_csv_opens_in_excel_with_hebrew(receipts, tmp_path):
    out = export_csv(receipts, tmp_path / "Receipts" / "k.csv")
    raw = out.read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")  # BOM → Excel detects UTF-8
    rows = list(csv.reader(raw.decode("utf-8-sig").splitlines()))
    assert rows[0][:3] == ["תאריך", "חברה", "סכום"]
    assert rows[1][1] == "Netflix" and rows[1][2] == "54.9"
    assert rows[1][7].replace("\\", "/") == "Netflix/2026/2026-10/2026-10-14_Netflix_54.90ILS.pdf"


def test_xlsx_has_three_rtl_sheets_numbers_links_and_totals(receipts, tmp_path):
    out = export_xlsx(receipts, tmp_path / "Receipts" / "k.xlsx")
    wb = load_workbook(out)
    assert wb.sheetnames == ["קבלות", "סיכום חודשי", "לפי חברה"]
    ws = wb["קבלות"]
    assert ws.sheet_view.rightToLeft is True
    assert ws["C2"].value == 54.9 and isinstance(ws["C2"].value, float)  # a number the accountant can SUM
    assert ws["H2"].hyperlink is not None
    assert ws.max_row == 6 and ws.auto_filter.ref == "A1:H6"

    monthly = wb["סיכום חודשי"]
    assert [c.value for c in monthly[1]] == ["חודש", "מספר קבלות", "ILS", "USD"]
    oct_row = next(row for row in monthly.iter_rows(min_row=2, values_only=True) if row[0] and "2026" in row[0] and "אוק" in row[0])
    assert oct_row[1] == 2 and oct_row[2] == pytest.approx(183.9)
    assert str(monthly.cell(row=monthly.max_row, column=2).value).startswith("=SUM(")

    companies = {row[0]: row for row in wb["לפי חברה"].iter_rows(min_row=2, values_only=True)}
    assert companies["Netflix"][1] == 2 and companies["Netflix"][2] == pytest.approx(109.8)


def test_xlsx_english_is_ltr(receipts, tmp_path):
    i18n.set_language("en")
    wb = load_workbook(export_xlsx(receipts, tmp_path / "e.xlsx"))
    assert wb.sheetnames[0] == "Receipts" and not wb["Receipts"].sheet_view.rightToLeft


def test_empty_export_still_valid(tmp_path):
    wb = load_workbook(export_xlsx([], tmp_path / "empty.xlsx"))
    assert wb["קבלות"].max_row == 1


# ---- background schedule ---------------------------------------------------------


def test_is_due():
    now = datetime(2026, 10, 20, 12, 0)
    assert is_due(None, now, 24)
    assert not is_due(datetime(2026, 10, 20, 1, 0), now, 24)
    assert is_due(datetime(2026, 10, 19, 11, 0), now, 24)
    assert is_due(datetime(2026, 10, 20, 5, 0), now, 6)


def test_run_at_login_round_trip(tmp_path, qapp):
    from PySide6.QtCore import QSettings

    s = QSettings(str(tmp_path / "run.ini"), QSettings.Format.IniFormat)
    assert not runs_at_login(s)
    set_run_at_login(True, s)
    assert runs_at_login(s) and "--background" in s.value("Kabala")
    set_run_at_login(False, s)
    assert not runs_at_login(s)
