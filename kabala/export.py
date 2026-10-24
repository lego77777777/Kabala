"""
Export for the accountant: Excel (receipts + monthly summary + by company) or CSV.

Excel files open right-to-left when the UI is Hebrew, amounts are real numbers (not text) so the
accountant can sum them, and each row links to its PDF.
"""

from __future__ import annotations

import csv
import os
from datetime import date, timedelta
from pathlib import Path

from kabala.i18n import is_rtl, month_label, t
from kabala.models import Receipt

PERIODS = ["all", "this_month", "last_month", "this_year", "last_year"]


def period_range(key: str, today: date | None = None) -> tuple[date, date] | None:
    """Inclusive (start, end) for a period key, or None for everything."""
    today = today or date.today()
    first_this_month = today.replace(day=1)
    if key == "this_month":
        return first_this_month, today
    if key == "last_month":
        end = first_this_month - timedelta(days=1)
        return end.replace(day=1), end
    if key == "this_year":
        return date(today.year, 1, 1), today
    if key == "last_year":
        return date(today.year - 1, 1, 1), date(today.year - 1, 12, 31)
    return None


def filter_period(receipts: list[Receipt], key: str, today: date | None = None) -> list[Receipt]:
    rng = period_range(key, today)
    rows = sorted(receipts, key=lambda r: (r.date, r.company))
    if rng is None:
        return rows
    start, end = rng[0].isoformat(), rng[1].isoformat()
    return [r for r in rows if start <= r.date <= end]


def _columns() -> list[str]:
    return list(t("export.columns"))


def _relative(path: str, base: Path) -> str:
    if not path:
        return ""
    try:
        return os.path.relpath(path, base)
    except ValueError:  # different drive on Windows
        return path


def _row(r: Receipt, base: Path) -> list:
    return [r.date, r.company, round(r.amount, 2), r.currency, r.invoice_number or "", r.subject, r.sender,
            _relative(r.file_path, base)]


def export_csv(receipts: list[Receipt], target: Path) -> Path:
    """UTF-8 with BOM, so Excel shows Hebrew correctly when the file is double-clicked."""
    target.parent.mkdir(parents=True, exist_ok=True)
    with open(target, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(_columns())
        for r in receipts:
            w.writerow(_row(r, target.parent))
    return target


def export_xlsx(receipts: list[Receipt], target: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    rtl = is_rtl()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="4F46E5")
    total_font = Font(bold=True)
    thin = Border(bottom=Side(style="thin", color="E4E4E7"))
    money_fmt = "#,##0.00"

    wb = Workbook()

    def sheet(title: str, headers: list[str], widths: list[int], first: bool = False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = title[:31]
        ws.sheet_view.rightToLeft = rtl
        ws.append(headers)
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
            cell = ws.cell(row=1, column=i)
            cell.font, cell.fill = header_font, header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 22
        ws.freeze_panes = "A2"
        return ws

    # 1) every receipt
    ws = sheet(t("export.sheet_receipts"), _columns(), [12, 22, 12, 9, 18, 40, 30, 50], first=True)
    for r in receipts:
        ws.append(_row(r, target.parent))
        row = ws.max_row
        ws.cell(row=row, column=3).number_format = money_fmt
        if r.file_path:
            link = ws.cell(row=row, column=8)
            link.hyperlink = _relative(r.file_path, target.parent)
            link.style = "Hyperlink"
        for c in range(1, 9):
            ws.cell(row=row, column=c).border = thin
    if receipts:
        ws.auto_filter.ref = f"A1:H{ws.max_row}"

    # 2) monthly summary, one column per currency
    currencies = sorted({r.currency for r in receipts}, key=lambda c: (c != "ILS", c))
    months: dict[str, dict[str, float]] = {}
    counts: dict[str, int] = {}
    for r in receipts:
        m = r.date[:7]
        months.setdefault(m, {})
        months[m][r.currency] = months[m].get(r.currency, 0.0) + r.amount
        counts[m] = counts.get(m, 0) + 1
    ws2 = sheet(t("export.sheet_monthly"), [t("export.month"), t("export.count"), *currencies], [16, 10, *[14] * len(currencies)])
    for m in sorted(months):
        ws2.append([month_label(m, long=True), counts[m], *[round(months[m].get(c, 0.0), 2) for c in currencies]])
    if months:
        last = ws2.max_row
        ws2.append([t("export.total"), f"=SUM(B2:B{last})",
                    *[f"=SUM({get_column_letter(3 + i)}2:{get_column_letter(3 + i)}{last})" for i in range(len(currencies))]])
        for c in range(1, 3 + len(currencies)):
            ws2.cell(row=ws2.max_row, column=c).font = total_font
    for col in range(3, 3 + len(currencies)):
        for row in range(2, ws2.max_row + 1):
            ws2.cell(row=row, column=col).number_format = money_fmt

    # 3) by company
    by_company: dict[tuple[str, str], tuple[int, float]] = {}
    for r in receipts:
        n, s = by_company.get((r.company, r.currency), (0, 0.0))
        by_company[(r.company, r.currency)] = (n + 1, s + r.amount)
    ws3 = sheet(t("export.sheet_companies"), [t("export.company"), t("export.count"), t("export.amount"), t("export.currency")],
                [26, 10, 14, 10])
    for (company, cur), (n, s) in sorted(by_company.items(), key=lambda kv: -kv[1][1]):
        ws3.append([company, n, round(s, 2), cur])
        ws3.cell(row=ws3.max_row, column=3).number_format = money_fmt

    target.parent.mkdir(parents=True, exist_ok=True)
    wb.save(target)
    return target
